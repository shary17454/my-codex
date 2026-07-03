from __future__ import annotations

import csv
import html
import json
import os
import sqlite3
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "y60_catalog_seed.json"
OUT = ROOT / "output"
DB_DIR = OUT / "database"
DOC_DIR = OUT / "docs"
PDF_DIR = OUT / "pdf"
MD_DIR = OUT / "markdown"

PART_COLUMNS = [
    "record_id",
    "plate_number",
    "plate_title_en",
    "plate_title_ar",
    "exploded_diagram_reference",
    "official_nissan_part_number",
    "english_name",
    "arabic_name",
    "english_description",
    "arabic_description",
    "quantity",
    "applicable_years",
    "applicable_vin",
    "applicable_market",
    "applicable_engine",
    "applicable_transmission",
    "applicable_trim",
    "applicable_body",
    "superseded_part_number",
    "replacement_part_number",
    "technical_notes_en",
    "technical_notes_ar",
    "source_status",
    "source_reference",
]

MISSING_EPC_COLUMNS = [
    "missing_record_id",
    "system_key",
    "system_en",
    "system_ar",
    "original_epc_plate_number",
    "original_plate_title",
    "arabic_plate_title",
    "missing_items",
    "reason_unavailable_en",
    "reason_unavailable_ar",
    "required_source_evidence",
    "status",
]


def ensure_dirs() -> None:
    for path in [DB_DIR, DOC_DIR, PDF_DIR, MD_DIR]:
        path.mkdir(parents=True, exist_ok=True)


def load_seed() -> dict:
    with DATA_PATH.open("r", encoding="utf-8") as f:
        return json.load(f)


def vr() -> str:
    return "VERIFICATION REQUIRED"


def build_records(seed: dict) -> dict:
    generated_at = datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    years = seed["years"]
    markets = seed["markets"]
    engines = seed["engines"]
    transmissions = seed["transmissions"]
    systems = seed["systems"]
    body_styles = seed["body_styles"]
    trims = seed["trims"]

    plates = []
    parts = []
    missing_epc_plates = []
    for idx, system in enumerate(systems, start=1):
        plate_number = f"Y60-{idx:03d}"
        plate = {
            "plate_number": plate_number,
            "system_key": system["key"],
            "title_en": f"{system['en']} - Master Verification Plate",
            "title_ar": f"{system['ar']} - لوحة تحقق رئيسية",
            "exploded_diagram_reference": vr(),
            "associated_parts_table": f"parts_{system['key']}",
            "source_status": "UNVERIFIED PLACEHOLDER",
            "source_reference": vr(),
        }
        plates.append(plate)
        parts.append(
            {
                "record_id": f"Y60-PART-{idx:04d}",
                "plate_number": plate_number,
                "plate_title_en": plate["title_en"],
                "plate_title_ar": plate["title_ar"],
                "exploded_diagram_reference": vr(),
                "official_nissan_part_number": vr(),
                "english_name": f"{system['en']} component - verification record",
                "arabic_name": f"قطعة من {system['ar']} - سجل تحقق",
                "english_description": (
                    "Official Nissan EPC part record not yet loaded. "
                    "Populate only from verified Nissan EPC plate, parts microfiche, FAST export, "
                    "or dealer supersession evidence."
                ),
                "arabic_description": (
                    "لم يتم تحميل سجل القطعة الرسمي من Nissan EPC بعد. "
                    "تتم التعبئة فقط من لوحة EPC موثقة أو ميكروفيش قطع أو تصدير FAST أو دليل بدائل من الوكيل."
                ),
                "quantity": vr(),
                "applicable_years": f"{years[0]}-{years[-1]} pending plate verification",
                "applicable_vin": vr(),
                "applicable_market": "All listed markets pending verification",
                "applicable_engine": "All listed engines pending verification",
                "applicable_transmission": "All listed transmissions pending verification",
                "applicable_trim": vr(),
                "applicable_body": vr(),
                "superseded_part_number": vr(),
                "replacement_part_number": vr(),
                "technical_notes_en": "Do not order, quote, or publish as a final part until verified against official Nissan data.",
                "technical_notes_ar": "لا تستخدم للطلب أو التسعير أو النشر النهائي قبل التحقق من بيانات نيسان الرسمية.",
                "source_status": "UNVERIFIED PLACEHOLDER",
                "source_reference": vr(),
            }
        )
        missing_epc_plates.append(
            {
                "missing_record_id": f"Y60-MISSING-EPC-{idx:04d}",
                "system_key": system["key"],
                "system_en": system["en"],
                "system_ar": system["ar"],
                "original_epc_plate_number": vr(),
                "original_plate_title": vr(),
                "arabic_plate_title": vr(),
                "missing_items": (
                    "Original EPC plate number; original plate title; complete exploded diagram; "
                    "all callout/reference numbers; every illustrated part; Nissan part numbers; "
                    "quantity; VIN, production date, engine, transmission, market applicability; "
                    "supersession and replacement data; technical notes."
                ),
                "reason_unavailable_en": (
                    "No official Nissan EPC/FAST/microfiche/source plate for this Y60 system is present "
                    "in the workspace. The original plate cannot be enumerated, reproduced, translated, "
                    "or indexed without source evidence."
                ),
                "reason_unavailable_ar": (
                    "لا توجد لوحة مصدر رسمية من Nissan EPC أو FAST أو الميكروفيش لهذا النظام من Y60 داخل مساحة العمل. "
                    "لا يمكن حصر اللوحة الأصلية أو إعادة إنتاجها أو ترجمتها أو فهرستها بدون دليل مصدر."
                ),
                "required_source_evidence": (
                    "Official Nissan EPC/FAST export, dealer EPC printout, parts microfiche scan, "
                    "or authorized original plate image/PDF with market/model applicability."
                ),
                "status": "MISSING - OFFICIAL EPC SOURCE REQUIRED",
            }
        )

    indexes = {
        "years": [{"year": year, "status": vr()} for year in years],
        "engines": engines,
        "markets": markets,
        "systems": systems,
        "transmissions": transmissions,
        "body_styles": body_styles,
        "trims": trims,
        "part_numbers": [{"official_nissan_part_number": vr(), "status": "No verified part numbers loaded"}],
    }

    return {
        "metadata": {
            **seed["catalog"],
            "generated_at": generated_at,
            "source_policy": "Verified official data only. Unknown data is marked VERIFICATION REQUIRED.",
        },
        "years": years,
        "markets": markets,
        "engines": engines,
        "transmissions": transmissions,
        "body_styles": body_styles,
        "trims": trims,
        "systems": systems,
        "plates": plates,
        "parts": parts,
        "missing_epc_plates": missing_epc_plates,
        "indexes": indexes,
        "quality_rules": seed["quality_rules"],
    }


def write_json(data: dict) -> None:
    with (DB_DIR / "nissan_patrol_y60_catalog.json").open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def write_csv(data: dict) -> None:
    with (DB_DIR / "nissan_patrol_y60_parts.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=PART_COLUMNS)
        writer.writeheader()
        writer.writerows(data["parts"])
    with (DB_DIR / "nissan_patrol_y60_plates.csv").open("w", encoding="utf-8-sig", newline="") as f:
        fields = list(data["plates"][0].keys())
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(data["plates"])
    with (DB_DIR / "nissan_patrol_y60_missing_epc_plates.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=MISSING_EPC_COLUMNS)
        writer.writeheader()
        writer.writerows(data["missing_epc_plates"])


def write_sqlite(data: dict) -> None:
    db = DB_DIR / "nissan_patrol_y60_catalog.sqlite"
    if db.exists():
        db.unlink()
    con = sqlite3.connect(db)
    cur = con.cursor()
    cur.execute(
        """
        CREATE TABLE parts (
            record_id TEXT PRIMARY KEY,
            plate_number TEXT,
            plate_title_en TEXT,
            plate_title_ar TEXT,
            exploded_diagram_reference TEXT,
            official_nissan_part_number TEXT,
            english_name TEXT,
            arabic_name TEXT,
            english_description TEXT,
            arabic_description TEXT,
            quantity TEXT,
            applicable_years TEXT,
            applicable_vin TEXT,
            applicable_market TEXT,
            applicable_engine TEXT,
            applicable_transmission TEXT,
            applicable_trim TEXT,
            applicable_body TEXT,
            superseded_part_number TEXT,
            replacement_part_number TEXT,
            technical_notes_en TEXT,
            technical_notes_ar TEXT,
            source_status TEXT,
            source_reference TEXT
        )
        """
    )
    cur.executemany(
        f"INSERT INTO parts ({','.join(PART_COLUMNS)}) VALUES ({','.join(['?'] * len(PART_COLUMNS))})",
        [[row[col] for col in PART_COLUMNS] for row in data["parts"]],
    )
    cur.execute(
        """
        CREATE TABLE plates (
            plate_number TEXT PRIMARY KEY,
            system_key TEXT,
            title_en TEXT,
            title_ar TEXT,
            exploded_diagram_reference TEXT,
            associated_parts_table TEXT,
            source_status TEXT,
            source_reference TEXT
        )
        """
    )
    plate_cols = list(data["plates"][0].keys())
    cur.executemany(
        f"INSERT INTO plates ({','.join(plate_cols)}) VALUES ({','.join(['?'] * len(plate_cols))})",
        [[row[col] for col in plate_cols] for row in data["plates"]],
    )
    cur.execute("CREATE INDEX idx_parts_part_number ON parts(official_nissan_part_number)")
    cur.execute("CREATE INDEX idx_parts_arabic_name ON parts(arabic_name)")
    cur.execute("CREATE INDEX idx_parts_english_name ON parts(english_name)")
    cur.execute("CREATE INDEX idx_parts_plate ON parts(plate_number)")
    cur.execute(
        """
        CREATE TABLE missing_epc_plates (
            missing_record_id TEXT PRIMARY KEY,
            system_key TEXT,
            system_en TEXT,
            system_ar TEXT,
            original_epc_plate_number TEXT,
            original_plate_title TEXT,
            arabic_plate_title TEXT,
            missing_items TEXT,
            reason_unavailable_en TEXT,
            reason_unavailable_ar TEXT,
            required_source_evidence TEXT,
            status TEXT
        )
        """
    )
    cur.executemany(
        f"INSERT INTO missing_epc_plates ({','.join(MISSING_EPC_COLUMNS)}) VALUES ({','.join(['?'] * len(MISSING_EPC_COLUMNS))})",
        [[row[col] for col in MISSING_EPC_COLUMNS] for row in data["missing_epc_plates"]],
    )
    cur.execute("CREATE INDEX idx_missing_epc_system ON missing_epc_plates(system_key)")
    con.commit()
    con.close()


def write_xlsx(data: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts"
    ws.append(PART_COLUMNS)
    for row in data["parts"]:
        ws.append([row[col] for col in PART_COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_cells in ws.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(length + 2, 14), 48)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for sheet_name, rows in [
        ("Plates", data["plates"]),
        ("Years", data["indexes"]["years"]),
        ("Engines", data["engines"]),
        ("Markets", data["markets"]),
        ("Systems", data["systems"]),
        ("Missing EPC Plates", data["missing_epc_plates"]),
        ("QC", qc_rows(data)),
    ]:
        sheet = wb.create_sheet(sheet_name)
        if rows:
            cols = list(rows[0].keys())
            sheet.append(cols)
            for row in rows:
                sheet.append([row.get(col, "") for col in cols])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for column_cells in sheet.columns:
                length = max(len(str(cell.value or "")) for cell in column_cells)
                sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(length + 2, 14), 52)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
    wb.save(DB_DIR / "nissan_patrol_y60_catalog.xlsx")


def qc_rows(data: dict) -> list[dict]:
    rows = []
    part_numbers = {}
    for part in data["parts"]:
        pn = part["official_nissan_part_number"]
        part_numbers[pn] = part_numbers.get(pn, 0) + 1
        checks = [
            ("missing_part_number", pn == vr()),
            ("missing_diagram", part["exploded_diagram_reference"] == vr()),
            ("missing_vin", part["applicable_vin"] == vr()),
            ("missing_source", part["source_reference"] == vr()),
            ("missing_arabic", not part["arabic_name"] or not part["arabic_description"]),
        ]
        for check_id, failed in checks:
            if failed:
                rows.append(
                    {
                        "severity": "BLOCKER",
                        "check_id": check_id,
                        "record_id": part["record_id"],
                        "message_en": f"{check_id} requires official Nissan source verification.",
                        "message_ar": f"{check_id} يتطلب التحقق من مصدر رسمي من نيسان.",
                    }
                )
    for missing in data["missing_epc_plates"]:
        rows.append(
            {
                "severity": "BLOCKER",
                "check_id": "missing_original_epc_plate",
                "record_id": missing["missing_record_id"],
                "message_en": "Original Nissan EPC plate is unavailable and must be supplied from an official source.",
                "message_ar": "لوحة Nissan EPC الأصلية غير متوفرة ويجب توفيرها من مصدر رسمي.",
            }
        )
    if any(count > 1 and pn != vr() for pn, count in part_numbers.items()):
        rows.append(
            {
                "severity": "WARNING",
                "check_id": "duplicate_part_number",
                "record_id": "",
                "message_en": "Duplicate verified part numbers found.",
                "message_ar": "تم العثور على أرقام قطع موثقة مكررة.",
            }
        )
    return rows


def write_qc(data: dict) -> None:
    rows = qc_rows(data)
    with (OUT / "quality_control_report.json").open("w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    with (OUT / "quality_control_report.csv").open("w", encoding="utf-8-sig", newline="") as f:
        fields = list(rows[0].keys()) if rows else ["severity", "check_id", "record_id", "message_en", "message_ar"]
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_markdown(data: dict) -> None:
    md = []
    meta = data["metadata"]
    md.append(f"# {meta['english_title']}")
    md.append(f"# {meta['arabic_title']}")
    md.append("")
    md.append("## Table of Contents")
    md.append("## فهرس المحتويات")
    md.extend(
        [
            "- [Scope / النطاق](#scope--النطاق)",
            "- [Data Integrity / سلامة البيانات](#data-integrity--سلامة-البيانات)",
            "- [Indexes / الفهارس](#indexes--الفهارس)",
            "- [Catalog Plates / لوحات الكتالوج](#catalog-plates--لوحات-الكتالوج)",
            "- [Parts Table / جدول القطع](#parts-table--جدول-القطع)",
            "- [Quality Control / ضبط الجودة](#quality-control--ضبط-الجودة)",
        ]
    )
    md.append("")
    md.append("## Scope / النطاق")
    md.append(meta["scope_note_en"])
    md.append(meta["scope_note_ar"])
    md.append("")
    md.append("## Data Integrity / سلامة البيانات")
    md.append(meta["data_integrity_en"])
    md.append(meta["data_integrity_ar"])
    md.append("")
    md.append("## Indexes / الفهارس")
    md.append(f"Years / السنوات: {', '.join(str(y) for y in data['years'])}")
    md.append("Engines / المحركات: " + ", ".join(engine["code"] for engine in data["engines"]))
    md.append("Markets / الأسواق: " + ", ".join(market["en"] for market in data["markets"]))
    md.append("")
    md.append("## Catalog Plates / لوحات الكتالوج")
    md.append("| Plate | English Title | Arabic Title | Diagram | Source |")
    md.append("|---|---|---|---|---|")
    for plate in data["plates"]:
        md.append(
            f"| {plate['plate_number']} | {plate['title_en']} | {plate['title_ar']} | "
            f"{plate['exploded_diagram_reference']} | {plate['source_status']} |"
        )
    md.append("")
    md.append("## Parts Table / جدول القطع")
    md.append("| Record | Plate | Part Number | English Name | Arabic Name | Status |")
    md.append("|---|---|---|---|---|---|")
    for part in data["parts"]:
        md.append(
            f"| {part['record_id']} | {part['plate_number']} | {part['official_nissan_part_number']} | "
            f"{part['english_name']} | {part['arabic_name']} | {part['source_status']} |"
        )
    md.append("")
    md.append("## Quality Control / ضبط الجودة")
    md.append(f"Open blocker count / عدد ملاحظات الحجب: {len(qc_rows(data))}")
    md.append("All blocker rows are intentional until official Nissan source records are imported.")
    md.append("كل ملاحظات الحجب مقصودة إلى أن يتم استيراد سجلات رسمية من نيسان.")

    (MD_DIR / "nissan_patrol_y60_catalog.md").write_text("\n".join(md), encoding="utf-8")


def doc_add_bilingual_heading(doc: Document, en: str, ar: str, level: int = 1) -> None:
    doc.add_heading(en, level=level)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(ar)
    run.bold = True
    run.font.name = "Arial"


def write_docx(data: dict) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)

    meta = data["metadata"]
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(meta["english_title"])
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(31, 78, 121)
    ar_title = doc.add_paragraph()
    ar_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ar_title.add_run(meta["arabic_title"])
    r.bold = True
    r.font.size = Pt(16)
    r.font.name = "Arial"

    doc.add_paragraph(meta["scope_note_en"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(meta["scope_note_ar"])
    doc.add_paragraph(meta["data_integrity_en"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(meta["data_integrity_ar"])

    doc_add_bilingual_heading(doc, "Indexes", "الفهارس", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "English"
    table.rows[0].cells[1].text = "العربية"
    for label_en, label_ar, value in [
        ("Years", "السنوات", ", ".join(str(y) for y in data["years"])),
        ("Engines", "المحركات", ", ".join(engine["code"] for engine in data["engines"])),
        ("Markets", "الأسواق", ", ".join(m["en"] for m in data["markets"])),
        ("Systems", "الأنظمة", ", ".join(s["en"] for s in data["systems"])),
    ]:
        row = table.add_row().cells
        row[0].text = f"{label_en}: {value}"
        row[1].text = f"{label_ar}: {value}"

    doc_add_bilingual_heading(doc, "Catalog Plates", "لوحات الكتالوج", 1)
    plate_table = doc.add_table(rows=1, cols=5)
    plate_table.style = "Table Grid"
    for i, header in enumerate(["Plate", "English Title", "Arabic Title", "Diagram", "Source"]):
        plate_table.rows[0].cells[i].text = header
    for plate in data["plates"]:
        row = plate_table.add_row().cells
        row[0].text = plate["plate_number"]
        row[1].text = plate["title_en"]
        row[2].text = plate["title_ar"]
        row[3].text = plate["exploded_diagram_reference"]
        row[4].text = plate["source_status"]

    doc_add_bilingual_heading(doc, "Parts Verification Table", "جدول التحقق من القطع", 1)
    part_table = doc.add_table(rows=1, cols=6)
    part_table.style = "Table Grid"
    for i, header in enumerate(["Record", "Plate", "Part Number", "English Name", "Arabic Name", "Status"]):
        part_table.rows[0].cells[i].text = header
    for part in data["parts"]:
        row = part_table.add_row().cells
        row[0].text = part["record_id"]
        row[1].text = part["plate_number"]
        row[2].text = part["official_nissan_part_number"]
        row[3].text = part["english_name"]
        row[4].text = part["arabic_name"]
        row[5].text = part["source_status"]

    doc_add_bilingual_heading(doc, "Quality Control", "ضبط الجودة", 1)
    doc.add_paragraph(f"Open blocker count: {len(qc_rows(data))}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f"عدد ملاحظات الحجب المفتوحة: {len(qc_rows(data))}")

    doc.save(DOC_DIR / "nissan_patrol_y60_bilingual_parts_catalog.docx")


def register_pdf_fonts() -> tuple[str, str]:
    regular = r"C:\Windows\Fonts\arial.ttf"
    bold = r"C:\Windows\Fonts\arialbd.ttf"
    if os.path.exists(regular):
        pdfmetrics.registerFont(TTFont("Arial", regular))
        if os.path.exists(bold):
            pdfmetrics.registerFont(TTFont("Arial-Bold", bold))
        return "Arial", "Arial-Bold"
    return "Helvetica", "Helvetica-Bold"


def write_pdf(data: dict) -> None:
    font, bold_font = register_pdf_fonts()
    path = PDF_DIR / "nissan_patrol_y60_bilingual_parts_catalog.reportlab_fallback.pdf"
    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleCustom", parent=styles["Title"], fontName=bold_font, fontSize=18, leading=22)
    h = ParagraphStyle("HeadingCustom", parent=styles["Heading2"], fontName=bold_font, fontSize=13, leading=16)
    body = ParagraphStyle("BodyCustom", parent=styles["BodyText"], fontName=font, fontSize=8.5, leading=11, alignment=TA_LEFT)
    cell = ParagraphStyle("CellCustom", parent=body, fontName=font, fontSize=6.5, leading=8)
    ar = ParagraphStyle("ArabicCustom", parent=body, fontName=font, alignment=TA_RIGHT)
    story = []
    meta = data["metadata"]
    story.append(Paragraph(meta["english_title"], title))
    story.append(Paragraph(meta["arabic_title"], ar))
    story.append(Spacer(1, 6))
    story.append(Paragraph(meta["scope_note_en"], body))
    story.append(Paragraph(meta["scope_note_ar"], ar))
    story.append(Paragraph(meta["data_integrity_en"], body))
    story.append(Paragraph(meta["data_integrity_ar"], ar))
    story.append(PageBreak())

    story.append(Paragraph("Catalog Plates", h))
    story.append(Paragraph("لوحات الكتالوج", ar))
    def c(value: str) -> Paragraph:
        return Paragraph(str(value), cell)

    plate_data = [[c("Plate"), c("English Title"), c("Arabic Title"), c("Diagram"), c("Source")]]
    for plate in data["plates"]:
        plate_data.append([
            c(plate["plate_number"]),
            c(plate["title_en"]),
            c(plate["title_ar"]),
            c(plate["exploded_diagram_reference"]),
            c(plate["source_status"]),
        ])
    t = Table(plate_data, colWidths=[25 * mm, 60 * mm, 70 * mm, 45 * mm, 45 * mm], repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("FONTNAME", (0, 0), (-1, 0), bold_font),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C7CED6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(t)
    story.append(PageBreak())

    story.append(Paragraph("Parts Verification Table", h))
    story.append(Paragraph("جدول التحقق من القطع", ar))
    part_data = [[c("Record"), c("Plate"), c("Part Number"), c("English Name"), c("Arabic Name"), c("Status")]]
    for part in data["parts"]:
        part_data.append([
            c(part["record_id"]),
            c(part["plate_number"]),
            c(part["official_nissan_part_number"]),
            c(part["english_name"]),
            c(part["arabic_name"]),
            c(part["source_status"]),
        ])
    t = Table(part_data, colWidths=[32 * mm, 24 * mm, 45 * mm, 65 * mm, 65 * mm, 45 * mm], repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("FONTNAME", (0, 0), (-1, 0), bold_font),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C7CED6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(t)
    doc.build(story)


def write_print_html(data: dict) -> None:
    meta = data["metadata"]

    def e(value: object) -> str:
        return html.escape(str(value))

    plate_rows = "\n".join(
        "<tr>"
        f"<td>{e(p['plate_number'])}</td>"
        f"<td>{e(p['title_en'])}<div class='ar'>{e(p['title_ar'])}</div></td>"
        f"<td>{e(p['exploded_diagram_reference'])}</td>"
        f"<td>{e(p['associated_parts_table'])}</td>"
        f"<td>{e(p['source_status'])}</td>"
        "</tr>"
        for p in data["plates"]
    )
    part_rows = "\n".join(
        "<tr>"
        f"<td>{e(p['record_id'])}</td>"
        f"<td>{e(p['plate_number'])}</td>"
        f"<td>{e(p['official_nissan_part_number'])}</td>"
        f"<td>{e(p['english_name'])}<div class='ar'>{e(p['arabic_name'])}</div></td>"
        f"<td>{e(p['quantity'])}</td>"
        f"<td>{e(p['source_status'])}</td>"
        "</tr>"
        for p in data["parts"]
    )
    qc_count = len(qc_rows(data))
    engines = ", ".join(e(engine["code"]) for engine in data["engines"])
    years = ", ".join(str(y) for y in data["years"])
    markets = ", ".join(e(m["en"]) for m in data["markets"])
    systems = ", ".join(e(s["en"]) for s in data["systems"])
    html_text = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>{e(meta['english_title'])}</title>
  <style>
    @page {{ size: A4 landscape; margin: 12mm; }}
    body {{ font-family: Arial, Tahoma, sans-serif; color: #111827; line-height: 1.35; }}
    h1 {{ color: #173f5f; font-size: 24px; margin: 0 0 8px; }}
    h2 {{ color: #173f5f; border-bottom: 1px solid #9fb3c8; padding-bottom: 4px; margin-top: 22px; }}
    .ar {{ direction: rtl; text-align: right; font-family: Tahoma, Arial, sans-serif; margin-top: 3px; }}
    .lead {{ max-width: 1000px; font-size: 12px; }}
    .toc a {{ color: #173f5f; text-decoration: none; }}
    table {{ width: 100%; border-collapse: collapse; table-layout: fixed; page-break-inside: auto; margin-top: 8px; }}
    th {{ background: #173f5f; color: #fff; text-align: left; font-size: 9px; padding: 5px; }}
    td {{ border: 1px solid #c7d0dc; vertical-align: top; font-size: 8.5px; padding: 4px; overflow-wrap: anywhere; }}
    tr {{ page-break-inside: avoid; page-break-after: auto; }}
    .note {{ border-left: 4px solid #b91c1c; background: #fff7ed; padding: 8px 10px; margin: 12px 0; font-size: 11px; }}
    .small {{ font-size: 10px; color: #374151; }}
    .page-break {{ page-break-before: always; }}
  </style>
</head>
<body>
  <h1>{e(meta['english_title'])}</h1>
  <h1 class="ar">{e(meta['arabic_title'])}</h1>
  <div class="lead">
    <p>{e(meta['scope_note_en'])}</p>
    <p class="ar">{e(meta['scope_note_ar'])}</p>
    <div class="note">
      <strong>Data Integrity</strong><br>
      {e(meta['data_integrity_en'])}
      <div class="ar"><strong>سلامة البيانات</strong><br>{e(meta['data_integrity_ar'])}</div>
    </div>
  </div>
  <h2>Hyperlinked Table of Contents</h2>
  <h2 class="ar">فهرس محتويات بروابط</h2>
  <p class="toc">
    <a href="#indexes">Indexes / الفهارس</a> |
    <a href="#plates">Catalog Plates / لوحات الكتالوج</a> |
    <a href="#parts">Parts Table / جدول القطع</a> |
    <a href="#qc">Quality Control / ضبط الجودة</a>
  </p>
  <h2 id="indexes">Indexes</h2>
  <h2 class="ar">الفهارس</h2>
  <p><strong>Years:</strong> {years}</p>
  <p><strong>Engines:</strong> {engines}</p>
  <p><strong>Markets:</strong> {markets}</p>
  <p><strong>Systems:</strong> {systems}</p>

  <div class="page-break"></div>
  <h2 id="plates">Catalog Plates</h2>
  <h2 class="ar">لوحات الكتالوج</h2>
  <table>
    <thead><tr><th style="width:9%">Plate</th><th style="width:34%">Title</th><th style="width:21%">Diagram</th><th style="width:20%">Associated Table</th><th style="width:16%">Source</th></tr></thead>
    <tbody>{plate_rows}</tbody>
  </table>

  <div class="page-break"></div>
  <h2 id="parts">Parts Verification Table</h2>
  <h2 class="ar">جدول التحقق من القطع</h2>
  <table>
    <thead><tr><th style="width:12%">Record</th><th style="width:8%">Plate</th><th style="width:18%">Part Number</th><th style="width:38%">Name</th><th style="width:8%">Qty</th><th style="width:16%">Status</th></tr></thead>
    <tbody>{part_rows}</tbody>
  </table>

  <h2 id="qc">Quality Control</h2>
  <h2 class="ar">ضبط الجودة</h2>
  <p>Open blocker count: {qc_count}. These blockers are intentional until official Nissan EPC data is imported.</p>
  <p class="ar">عدد ملاحظات الحجب المفتوحة: {qc_count}. هذه الملاحظات مقصودة إلى أن يتم استيراد بيانات Nissan EPC الرسمية.</p>
  <p class="small">Generated from canonical JSON. Unknown official data remains marked VERIFICATION REQUIRED.</p>
</body>
</html>
"""
    (PDF_DIR / "nissan_patrol_y60_bilingual_parts_catalog.html").write_text(html_text, encoding="utf-8")


def write_missing_epc_reports(data: dict) -> None:
    md = [
        "# Missing EPC Plates Report",
        "# تقرير لوحات EPC المفقودة",
        "",
        "This report is mandatory because no official Nissan EPC source plates are present in the workspace.",
        "هذا التقرير إلزامي لأن مساحة العمل لا تحتوي على لوحات مصدر رسمية من Nissan EPC.",
        "",
        "| Missing Record | System | Original EPC Plate Number | Original Plate Title | Missing Items | Status |",
        "|---|---|---|---|---|---|",
    ]
    for missing in data["missing_epc_plates"]:
        md.append(
            f"| {missing['missing_record_id']} | {missing['system_en']} / {missing['system_ar']} | "
            f"{missing['original_epc_plate_number']} | {missing['original_plate_title']} | "
            f"{missing['missing_items']} | {missing['status']} |"
        )
    (MD_DIR / "missing_epc_plates_report.md").write_text("\n".join(md), encoding="utf-8")

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    doc.add_heading("Missing EPC Plates Report", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("تقرير لوحات EPC المفقودة").bold = True
    doc.add_paragraph("This report is mandatory because no official Nissan EPC source plates are present in the workspace.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("هذا التقرير إلزامي لأن مساحة العمل لا تحتوي على لوحات مصدر رسمية من Nissan EPC.")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for i, header in enumerate(["Missing Record", "System", "Original EPC Plate", "Original Title", "Missing Items", "Status"]):
        table.rows[0].cells[i].text = header
    for missing in data["missing_epc_plates"]:
        row = table.add_row().cells
        row[0].text = missing["missing_record_id"]
        row[1].text = f"{missing['system_en']} / {missing['system_ar']}"
        row[2].text = missing["original_epc_plate_number"]
        row[3].text = missing["original_plate_title"]
        row[4].text = missing["missing_items"]
        row[5].text = missing["status"]
    doc.save(DOC_DIR / "missing_epc_plates_report.docx")

    def e(value: object) -> str:
        return html.escape(str(value))

    rows = "\n".join(
        "<tr>"
        f"<td>{e(m['missing_record_id'])}</td>"
        f"<td>{e(m['system_en'])}<div class='ar'>{e(m['system_ar'])}</div></td>"
        f"<td>{e(m['original_epc_plate_number'])}</td>"
        f"<td>{e(m['original_plate_title'])}<div class='ar'>{e(m['arabic_plate_title'])}</div></td>"
        f"<td>{e(m['missing_items'])}</td>"
        f"<td>{e(m['status'])}</td>"
        "</tr>"
        for m in data["missing_epc_plates"]
    )
    html_text = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>Missing EPC Plates Report</title>
  <style>
    @page {{ size: A4 landscape; margin: 12mm; }}
    body {{ font-family: Arial, Tahoma, sans-serif; color: #111827; line-height: 1.35; }}
    h1 {{ color: #173f5f; font-size: 24px; margin: 0 0 8px; }}
    .ar {{ direction: rtl; text-align: right; font-family: Tahoma, Arial, sans-serif; }}
    .note {{ border-left: 4px solid #b91c1c; background: #fff7ed; padding: 8px 10px; margin: 12px 0; font-size: 11px; }}
    table {{ width: 100%; border-collapse: collapse; table-layout: fixed; margin-top: 12px; }}
    th {{ background: #173f5f; color: white; padding: 5px; font-size: 9px; text-align: left; }}
    td {{ border: 1px solid #c7d0dc; padding: 4px; font-size: 8.5px; vertical-align: top; overflow-wrap: anywhere; }}
    tr {{ page-break-inside: avoid; }}
  </style>
</head>
<body>
  <h1>Missing EPC Plates Report</h1>
  <h1 class="ar">تقرير لوحات EPC المفقودة</h1>
  <div class="note">
    Official Nissan EPC plates, exploded diagrams, callouts, and part numbers are unavailable in this workspace. They must be supplied from official source evidence before a complete original EPC catalog can be generated.
    <div class="ar">لوحات Nissan EPC الأصلية والرسومات التفجيرية وأرقام النداء وأرقام القطع غير متوفرة داخل مساحة العمل. يجب توفيرها من دليل مصدر رسمي قبل إنشاء كتالوج EPC أصلي كامل.</div>
  </div>
  <table>
    <thead><tr><th style="width:12%">Missing Record</th><th style="width:15%">System</th><th style="width:15%">Original EPC Plate</th><th style="width:16%">Original Title</th><th style="width:30%">Missing Items</th><th style="width:12%">Status</th></tr></thead>
    <tbody>{rows}</tbody>
  </table>
</body>
</html>
"""
    (PDF_DIR / "missing_epc_plates_report.html").write_text(html_text, encoding="utf-8")


def write_schema_docs() -> None:
    schema = {
        "parts_table_columns": PART_COLUMNS,
        "missing_epc_plate_columns": MISSING_EPC_COLUMNS,
        "required_policy": {
            "official_nissan_part_number": "Must be sourced from official Nissan EPC/FAST/dealer source or marked VERIFICATION REQUIRED.",
            "arabic_fields": "Must use professional GCC/Saudi workshop terminology.",
            "model_codes": "Do not translate engine, transmission, model, axle, VIN, or part-number codes.",
        },
    }
    (DB_DIR / "schema.json").write_text(json.dumps(schema, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# Catalog Schema", "", "## Parts Columns", ""]
    lines.extend(f"- `{col}`" for col in PART_COLUMNS)
    lines.extend(["", "## Missing EPC Plate Columns", ""])
    lines.extend(f"- `{col}`" for col in MISSING_EPC_COLUMNS)
    (MD_DIR / "schema.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    ensure_dirs()
    seed = load_seed()
    data = build_records(seed)
    write_json(data)
    write_csv(data)
    write_sqlite(data)
    write_xlsx(data)
    write_markdown(data)
    write_docx(data)
    write_print_html(data)
    write_missing_epc_reports(data)
    write_qc(data)
    write_schema_docs()
    print(json.dumps({
        "parts": len(data["parts"]),
        "plates": len(data["plates"]),
        "missing_epc_plates": len(data["missing_epc_plates"]),
        "qc_blockers": len(qc_rows(data)),
        "output": str(OUT),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
